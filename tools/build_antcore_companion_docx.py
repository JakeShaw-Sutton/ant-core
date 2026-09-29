#!/usr/bin/env python3
"""Build the Ant Core companion Word document from README.md."""

from __future__ import annotations

import re
import tempfile
from pathlib import Path
from typing import Iterable

from PIL import Image
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
README = ROOT / "README.md"
OUTPUT = ROOT / "docs" / "Ant Core Companion Guide.docx"

BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
MUTED = RGBColor(0x55, 0x55, 0x55)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
BORDER = "A6B6C8"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def mark_first_row_as_header(table) -> None:
    if not table.rows:
        return
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def set_table_geometry(table, widths: list[float]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.allow_autofit = False
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    dxa_widths: list[int] = []
    for index, width in enumerate(widths):
        if index == len(widths) - 1:
            dxa_widths.append(9360 - sum(dxa_widths))
        else:
            dxa_widths.append(round(width * 1440))

    for dxa_width in dxa_widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(dxa_width))
        grid.append(col)

    for row in table.rows:
        for cell, width, dxa_width in zip(row.cells, widths, dxa_widths):
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(dxa_width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_font(run, name="Calibri", size: float | None = None, color: RGBColor | None = None, bold: bool | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold


def add_runs_from_inline_markdown(paragraph, text: str, base_size: float = 11) -> None:
    text = text.replace("\\|", "|")
    pattern = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            set_font(run, size=base_size)
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_font(run, name="Consolas", size=max(8.5, base_size - 1), color=DARK_BLUE)
        else:
            run = paragraph.add_run(token[2:-2])
            set_font(run, size=base_size, bold=True)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_font(run, size=base_size)


def clean_link_text(text: str) -> str:
    def repl(match: re.Match[str]) -> str:
        label, target = match.group(1), match.group(2)
        return f"{label} ({target})"

    return re.sub(r"\[([^\]]+)\]\(([^)]+)\)", repl, text)


def add_paragraph(doc: Document, text: str, style: str | None = None) -> None:
    paragraph = doc.add_paragraph(style=style)
    add_runs_from_inline_markdown(paragraph, clean_link_text(text))
    if style in {"List Bullet", "List Number"}:
        paragraph.paragraph_format.left_indent = Inches(0.375)
        paragraph.paragraph_format.first_line_indent = Inches(-0.188)
        paragraph.paragraph_format.space_after = Pt(4)


def add_code_block(doc: Document, lines: Iterable[str]) -> None:
    for line in lines:
        paragraph = doc.add_paragraph(style="Code Block")
        run = paragraph.add_run(line if line else " ")
        set_font(run, name="Consolas", size=8.5, color=DARK_BLUE)


def parse_table(lines: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in lines:
        stripped = line.strip()
        if re.fullmatch(r"\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?", stripped):
            continue
        if "|" not in stripped:
            continue
        if stripped.startswith("|"):
            stripped = stripped[1:]
        if stripped.endswith("|"):
            stripped = stripped[:-1]
        rows.append([cell.strip() for cell in stripped.split("|")])
    return rows


def column_widths(col_count: int) -> list[float]:
    if col_count <= 1:
        return [6.5]
    if col_count == 2:
        return [1.9, 4.6]
    if col_count == 3:
        return [1.45, 2.35, 2.7]
    if col_count == 4:
        return [1.3, 1.55, 1.8, 1.85]
    return [6.5 / col_count] * col_count


def add_markdown_table(doc: Document, lines: list[str]) -> None:
    rows = parse_table(lines)
    if not rows:
        return
    col_count = max(len(row) for row in rows)
    for row in rows:
        row.extend([""] * (col_count - len(row)))
    widths = column_widths(col_count)
    table = doc.add_table(rows=len(rows), cols=col_count)
    set_table_geometry(table, widths)
    set_table_borders(table)
    mark_first_row_as_header(table)
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_runs_from_inline_markdown(paragraph, clean_link_text(value), base_size=9.5)
            if r_idx == 0:
                set_cell_shading(cell, LIGHT_BLUE)
                for run in paragraph.runs:
                    run.bold = True
            elif col_count <= 3 and c_idx == 0:
                set_cell_shading(cell, LIGHT_GRAY)
                for run in paragraph.runs:
                    run.bold = True
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)


def split_tall_image(source: Path, out_dir: Path, max_ratio: float = 1.35) -> list[Path]:
    with Image.open(source) as image:
        width, height = image.size
        if height / max(width, 1) <= max_ratio:
            return [source]
        chunk_height = int(width * max_ratio)
        chunks = []
        for index, top in enumerate(range(0, height, chunk_height)):
            bottom = min(height, top + chunk_height)
            if bottom - top < 240 and chunks:
                prev = chunks.pop()
                with Image.open(prev) as prev_img:
                    merged = Image.new("RGB", (width, prev_img.height + bottom - top), (255, 255, 255))
                    merged.paste(prev_img.convert("RGB"), (0, 0))
                    merged.paste(image.crop((0, top, width, bottom)).convert("RGB"), (0, prev_img.height))
                    prev.unlink()
                    merged.save(prev)
                    chunks.append(prev)
                continue
            chunk = out_dir / f"{source.stem}-part-{index + 1}.png"
            image.crop((0, top, width, bottom)).save(chunk)
            chunks.append(chunk)
        return chunks


def add_image_figure(doc: Document, image_path: Path, alt: str, temp_dir: Path) -> None:
    if not image_path.exists():
        add_paragraph(doc, f"Screenshot missing: {image_path}")
        return
    chunks = split_tall_image(image_path, temp_dir)
    multi = len(chunks) > 1
    for index, chunk in enumerate(chunks):
        caption = alt if not multi else f"{alt} (part {index + 1} of {len(chunks)})"
        paragraph = doc.add_paragraph(style="Caption")
        run = paragraph.add_run(caption)
        set_font(run, size=9, color=MUTED, bold=True)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic = doc.add_paragraph()
        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        with Image.open(chunk) as image:
            width, height = image.size
            display_width = 6.25
            display_height = display_width * height / width
            if display_height > 7.35:
                display_height = 7.35
                display_width = display_height * width / height
        shape = pic.add_run().add_picture(str(chunk), width=Inches(display_width))
        shape._inline.docPr.set("title", caption)
        shape._inline.docPr.set("descr", caption)
        after = doc.add_paragraph()
        after.paragraph_format.space_after = Pt(6)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_style in ("List Bullet", "List Number"):
        style = styles[list_style]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code.font.size = Pt(8.5)
    code.font.color.rgb = DARK_BLUE
    code.paragraph_format.left_indent = Inches(0.15)
    code.paragraph_format.right_indent = Inches(0.15)
    code.paragraph_format.space_before = Pt(0)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.line_spacing = 1.0

    caption = styles["Caption"]
    caption.font.name = "Calibri"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    caption.font.size = Pt(9)
    caption.font.italic = False
    caption.font.color.rgb = MUTED
    caption.paragraph_format.space_before = Pt(6)
    caption.paragraph_format.space_after = Pt(3)

    header = section.header.paragraphs[0]
    header.text = "Ant Core Companion Guide"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        set_font(run, size=9, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.text = "Generated from README.md"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        set_font(run, size=9, color=MUTED)


def add_cover(doc: Document) -> None:
    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(80)
    title.paragraph_format.space_after = Pt(8)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Ant Core")
    set_font(run, size=30, color=BLUE, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Companion Guide")
    set_font(run, size=18, color=DARK_BLUE, bold=True)

    deck = doc.add_paragraph()
    deck.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = deck.add_run("Overview, setup, safety model, feature walkthrough, screenshots, and validation tools for the Seeed XIAO ESP32-S3 Sense antweight robot controller.")
    set_font(run, size=11, color=MUTED)

    note_table = doc.add_table(rows=4, cols=2)
    set_table_geometry(note_table, [1.8, 4.7])
    set_table_borders(note_table)
    mark_first_row_as_header(note_table)
    rows = [
        ("Source", "README.md"),
        ("Preset", "compact_reference_guide"),
        ("Target", "Word / DOCX companion manual"),
        ("Screenshot note", "Screenshots are captures of the shipped Ant Core web UI and support tools."),
    ]
    for idx, (label, value) in enumerate(rows):
        note_table.cell(idx, 0).text = label
        note_table.cell(idx, 1).text = value
        set_cell_shading(note_table.cell(idx, 0), LIGHT_BLUE)
        for cell in note_table.row_cells(idx):
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_font(run, size=10, bold=(cell is note_table.cell(idx, 0)))

    doc.add_page_break()


def build() -> None:
    markdown = README.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_document(doc)
    add_cover(doc)

    table_buffer: list[str] = []
    code_buffer: list[str] = []
    in_code = False

    def flush_table() -> None:
        nonlocal table_buffer
        if table_buffer:
            add_markdown_table(doc, table_buffer)
            table_buffer = []

    def flush_code() -> None:
        nonlocal code_buffer
        if code_buffer:
            add_code_block(doc, code_buffer)
            code_buffer = []

    with tempfile.TemporaryDirectory(prefix="antcore-docx-assets-") as tmp:
        temp_dir = Path(tmp)
        for raw in markdown:
            line = raw.rstrip()
            if line.startswith("```"):
                if in_code:
                    flush_code()
                    in_code = False
                else:
                    flush_table()
                    in_code = True
                continue
            if in_code:
                code_buffer.append(line)
                continue

            if not line.strip():
                flush_table()
                continue

            if re.match(r"^\|.*\|$", line.strip()):
                table_buffer.append(line)
                continue
            flush_table()

            image_match = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line.strip())
            if image_match:
                alt, target = image_match.groups()
                add_image_figure(doc, ROOT / target, alt or Path(target).stem, temp_dir)
                continue

            if line.startswith("# "):
                # Cover already supplies the main title.
                continue
            if line.startswith("#### "):
                doc.add_heading(clean_link_text(line[5:].strip()), level=3)
                continue
            if line.startswith("## "):
                doc.add_heading(clean_link_text(line[3:].strip()), level=1)
                continue
            if line.startswith("### "):
                doc.add_heading(clean_link_text(line[4:].strip()), level=2)
                continue

            if line.startswith("- "):
                add_paragraph(doc, line[2:].strip(), style="List Bullet")
                continue
            numbered = re.match(r"^\d+\.\s+(.*)$", line)
            if numbered:
                add_paragraph(doc, numbered.group(1).strip(), style="List Number")
                continue

            add_paragraph(doc, line.strip())

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
